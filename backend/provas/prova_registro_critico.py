"""Prova da Fase 24: registro critico para de mentir, e o export deixa rastro.

Duas coisas diferentes, e a primeira vem antes da segunda de proposito.

O CONSERTO: os dois eventos de registro critico estavam TROCADOS no codigo.
`routes/tarefas.py` CRIAVA uma excecao de obrigacao e emitia
`EXCLUSAO_REGISTRO_CRITICO`; `routes/obrigacoes.py` APAGAVA a excecao e emitia
`EDICAO_REGISTRO_CRITICO`. Nao muda comportamento nenhum, e por isso nenhuma
prova pegou: o que quebra e o FILTRO, que e a unica razao de o nome ser
padronizado. Quem procurava exclusao achava uma criacao.

O ACRESCIMO: `CRIACAO_REGISTRO_CRITICO` e `EXPORT_DADOS` nao existiam. Entram
nas quatro tabelas que o usuario definiu como criticas (usuario, permissao,
empresa e obrigacao), e nao nas 64 rotas de mutacao do app.

    cd backend && ./venv/bin/python provas/prova_registro_critico.py

Mede o que sai no stdout de verdade, num SQLite temporario.
"""

from __future__ import annotations  # producao e 3.12, a maquina local e 3.9

import io
import json
import os
import sys
import tempfile
from contextlib import redirect_stdout
from pathlib import Path

_tmp = tempfile.mkdtemp(prefix="prova_critico_")
_banco = os.path.join(_tmp, "prova.db")

os.environ["DATABASE_URL"] = f"sqlite:///{_banco}"
os.environ["SECRET_KEY"] = "chave-jwt-de-teste"
os.environ["ZOARIA_SSO_SECRET"] = ""

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient                      # noqa: E402

from app.auth import get_password_hash                         # noqa: E402
from app.database import Base, engine, SessionLocal            # noqa: E402
from app.models import (Empresa, Obrigacao, Setor, Tarefa,     # noqa: E402
                        Usuario)
from app.main import app                                       # noqa: E402

Base.metadata.create_all(bind=engine)

SENHA = "senha-boa-123"
EMAIL_ADMIN = "chefe@bps4.com.br"

CAMPOS = ["timestamp", "level", "event", "user_id", "ip", "request_id",
          "path", "method"]

PROIBIDAS = ["senha", "password", "token", "authorization", "cookie", "csrf",
             "cpf", "cartao", "secret"]

# Os tres nomes da tabela da nota. Nenhum outro nome de evento de cadastro pode
# sair do app: `event` fora da tabela e o anti-padrao "nao da pra filtrar".
CRIACAO = "CRIACAO_REGISTRO_CRITICO"
EDICAO = "EDICAO_REGISTRO_CRITICO"
EXCLUSAO = "EXCLUSAO_REGISTRO_CRITICO"

cliente = TestClient(app)
# Para os itens em que a rota estoura de proposito: sem isto o TestClient
# relanca a excecao em vez de devolver a resposta que o cliente real veria.
cliente_erro = TestClient(app, raise_server_exceptions=False)
ids = {}


def semear():
    db = SessionLocal()
    try:
        admin = Usuario(nome="Chefe", email=EMAIL_ADMIN, grupo="admin",
                        senha_hash=get_password_hash(SENHA), ativo=True)
        empresa = Empresa(razao_social="Cliente da Excecao",
                          cnpj="00000000000191")
        setor = Setor(nome="Fiscal")
        db.add_all([admin, empresa, setor])
        db.commit()
        obrigacao = Obrigacao(nome="DAS mensal", setor_id=setor.id)
        db.add(obrigacao)
        db.commit()
        tarefa = Tarefa(titulo="DAS de setembro", empresa_id=empresa.id,
                        setor_id=setor.id, responsavel_id=admin.id,
                        obrigacao_id=obrigacao.id)
        db.add(tarefa)
        db.commit()
        ids.update(admin=admin.id, empresa=empresa.id, setor=setor.id,
                   obrigacao=obrigacao.id, tarefa=tarefa.id)
    finally:
        db.close()


def entrar(email):
    r = cliente.post("/api/auth/login", json={"email": email, "senha": SENHA})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def capturar(funcao):
    buffer = io.StringIO()
    with redirect_stdout(buffer):
        retorno = funcao()
    linhas = []
    for bruta in buffer.getvalue().splitlines():
        bruta = bruta.strip()
        if not bruta.startswith("{"):
            continue
        try:
            linhas.append(json.loads(bruta))
        except json.JSONDecodeError:
            continue
    return retorno, linhas


def eventos(linhas, nome):
    return [linha for linha in linhas if linha.get("event") == nome]


def evento(linhas, nome):
    achados = eventos(linhas, nome)
    return achados[0] if achados else None


semear()
falhou = []
todas_as_linhas = []


def checa(n, descricao, condicao):
    if condicao:
        print(f"  ok  {n:>3}. {descricao}")
    else:
        print(f"FALHA  {n:>3}. {descricao}")
        falhou.append(n)


admin = entrar(EMAIL_ADMIN)


def registra(funcao):
    """Roda, guarda as linhas no acervo geral e devolve (resposta, linhas)."""
    resposta, linhas = capturar(funcao)
    todas_as_linhas.extend(linhas)
    return resposta, linhas


# ------------------------------------- (a) o conserto dos dois nomes trocados
r_excecao, l_excecao = registra(
    lambda: cliente.post(f"/api/tarefas/{ids['tarefa']}/nao-se-aplica",
                         json={"motivo": "cliente nao e optante"},
                         headers=admin))

checa(1, "criar excecao de obrigacao emite CRIACAO, e nao EXCLUSAO",
      r_excecao.status_code == 200
      and evento(l_excecao, CRIACAO) is not None
      and evento(l_excecao, EXCLUSAO) is None)
e_criacao = evento(l_excecao, CRIACAO)
checa(2, "e a linha diz a tabela e o alvo",
      e_criacao is not None
      and e_criacao.get("tabela") == "obrigacao_excecao"
      and e_criacao.get("empresa_id") == ids["empresa"])

# O id da excecao vem da lista do cadastro da obrigacao.
excecoes = cliente.get(f"/api/obrigacoes/{ids['obrigacao']}/excecoes",
                       headers=admin)
excecao_id = (excecoes.json()[0]["id"] if excecoes.status_code == 200
              and excecoes.json() else None)

r_desfaz, l_desfaz = registra(
    lambda: cliente.delete(
        f"/api/obrigacoes/{ids['obrigacao']}/excecoes/{excecao_id}",
        headers=admin))

checa(3, "apagar a excecao emite EXCLUSAO, e nao EDICAO",
      r_desfaz.status_code == 200
      and evento(l_desfaz, EXCLUSAO) is not None
      and evento(l_desfaz, EDICAO) is None)


# ------------------------------------------ (b) usuario: criar, editar, apagar
r_novo, l_novo = registra(
    lambda: cliente.post("/api/usuarios",
                         json={"nome": "Recem Chegado",
                               "email": "novo@bps4.com.br",
                               "senha": SENHA, "grupo": "analista"},
                         headers=admin))
novo_id = r_novo.json().get("id") if r_novo.status_code == 201 else None
e_novo = evento(l_novo, CRIACAO)
checa(4, "criar usuario emite CRIACAO com a tabela e o id do alvo",
      e_novo is not None and e_novo.get("tabela") == "usuario"
      and e_novo.get("alvo_id") == novo_id)

_, l_edita = registra(
    lambda: cliente.put(f"/api/usuarios/{novo_id}",
                        json={"cargo": "Auxiliar"}, headers=admin))
checa(5, "editar usuario emite EDICAO",
      evento(l_edita, EDICAO) is not None)

r_apaga, l_apaga = registra(
    lambda: cliente.delete(f"/api/usuarios/{novo_id}", headers=admin))
e_apaga = evento(l_apaga, EXCLUSAO)
checa(6, "apagar usuario emite EXCLUSAO",
      r_apaga.status_code == 200 and e_apaga is not None
      and e_apaga.get("tabela") == "usuario")
# A nota lista soft delete DENTRO do evento de exclusao ("Idem (soft delete)"),
# entao os dois desfechos usam o mesmo nome, e um campo distingue. Sem isso,
# quem le o log nao sabe se o cadastro sumiu ou so foi inativado.
checa(7, "e a linha diz se foi exclusao de fato ou apenas inativacao",
      e_apaga is not None and "inativado" in e_apaga)


# ------------------------------------------- (c) empresa: criar e apagar
r_emp, l_emp = registra(
    lambda: cliente.post("/api/empresas",
                         json={"razao_social": "Empresa Nova",
                               "cnpj": "11222333000181"},
                         headers=admin))
emp_id = r_emp.json().get("id") if r_emp.status_code == 201 else None
checa(8, "criar empresa emite CRIACAO com a tabela",
      evento(l_emp, CRIACAO) is not None
      and evento(l_emp, CRIACAO).get("tabela") == "empresa")

_, l_emp_edita = registra(
    lambda: cliente.put(f"/api/empresas/{emp_id}",
                        json={"razao_social": "Empresa Renomeada"},
                        headers=admin))
checa(9, "editar empresa emite EDICAO",
      evento(l_emp_edita, EDICAO) is not None)

_, l_emp_apaga = registra(
    lambda: cliente.delete(f"/api/empresas/{emp_id}", headers=admin))
checa(10, "apagar empresa emite EXCLUSAO",
      evento(l_emp_apaga, EXCLUSAO) is not None)


# ---------------------------------------- (d) obrigacao: criar, editar, apagar
r_obr, l_obr = registra(
    lambda: cliente.post("/api/obrigacoes",
                         json={"nome": "EFD Contribuicoes",
                               "setor_id": ids["setor"]},
                         headers=admin))
obr_id = r_obr.json().get("id") if r_obr.status_code == 201 else None
checa(11, "criar obrigacao emite CRIACAO com a tabela",
      evento(l_obr, CRIACAO) is not None
      and evento(l_obr, CRIACAO).get("tabela") == "obrigacao")

_, l_obr_edita = registra(
    lambda: cliente.put(f"/api/obrigacoes/{obr_id}",
                        json={"nome": "EFD Contribuicoes mensal"},
                        headers=admin))
checa(12, "editar obrigacao emite EDICAO",
      evento(l_obr_edita, EDICAO) is not None)

_, l_obr_apaga = registra(
    lambda: cliente.delete(f"/api/obrigacoes/{obr_id}?definitivo=true",
                           headers=admin))
checa(13, "apagar obrigacao emite EXCLUSAO",
      evento(l_obr_apaga, EXCLUSAO) is not None)


# ------------------------------------------------- (e) o documento excluido
# `DOCUMENTO_EXCLUIDO` era nome fora da tabela da nota, e por isso nao aparecia
# em nenhum filtro de exclusao. Vira o nome padronizado, e o antigo vira campo.
db = SessionLocal()
t = db.query(Tarefa).filter(Tarefa.id == ids["tarefa"]).first()
t.anexo_nome = "comprovante-de-teste.pdf"
db.commit()
db.close()

_, l_doc = registra(
    lambda: cliente.delete(f"/api/tarefas/{ids['tarefa']}/documento",
                           headers=admin))
e_doc = evento(l_doc, EXCLUSAO)
checa(14, "excluir documento emite EXCLUSAO, e nao um nome proprio",
      e_doc is not None and evento(l_doc, "DOCUMENTO_EXCLUIDO") is None)
checa(15, "e a tabela diz que foi anexo de tarefa",
      e_doc is not None and e_doc.get("tabela") == "tarefa_anexo")


# ------------------------------------------------------------- (f) os lotes
def _planilha_empresas():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["razao social", "cnpj"])
    for i, cnpj in enumerate(["11444777000161", "34028316000103"]):
        ws.append([f"Importada {i}", cnpj])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_, l_import = registra(
    lambda: cliente.post(
        "/api/empresas/importar",
        files={"arquivo": ("empresas.xlsx", _planilha_empresas(),
                           "application/vnd.openxmlformats-officedocument"
                           ".spreadsheetml.sheet")},
        headers=admin))
criacoes_no_lote = eventos(l_import, CRIACAO)
# Importar 200 empresas viraria 200 linhas e afogaria o log que este trabalho
# existe para tornar legivel. Uma linha, com a contagem.
checa(16, f"importar empresas emite UMA linha, nao uma por registro "
          f"(saiu: {len(criacoes_no_lote)})",
      len(criacoes_no_lote) == 1)
checa(17, "e a linha do lote traz a contagem do que entrou",
      bool(criacoes_no_lote)
      and isinstance(criacoes_no_lote[0].get("quantidade"), int)
      and criacoes_no_lote[0].get("lote") is True)

db = SessionLocal()
alvos = [o.id for o in db.query(Obrigacao).all()]
db.close()
_, l_lote = registra(
    lambda: cliente.post("/api/obrigacoes/excluir-lote",
                         json={"ids": alvos, "definitivo": True},
                         headers=admin))
exclusoes_no_lote = eventos(l_lote, EXCLUSAO)
checa(18, f"excluir em lote emite UMA linha com a contagem "
          f"(saiu: {len(exclusoes_no_lote)})",
      len(exclusoes_no_lote) == 1
      and exclusoes_no_lote[0].get("lote") is True)


# ------------------------------------------------------- (g) EXPORT_DADOS
# Duas obrigacoes de proposito, e depois do lote que esvaziou o cadastro: sem
# elas o relatorio sairia com zero linhas, e o item da contagem passaria sem
# medir contagem nenhuma. Foi o que aconteceu na primeira rodada.
registra(lambda: cliente.post("/api/obrigacoes",
                              json={"nome": "DCTFWeb", "setor_id": ids["setor"]},
                              headers=admin))
registra(lambda: cliente.post("/api/obrigacoes",
                              json={"nome": "ECD", "setor_id": ids["setor"]},
                              headers=admin))

_, l_export = registra(
    lambda: cliente.get("/api/obrigacoes/relatorio", headers=admin))
e_export = evento(l_export, "EXPORT_DADOS")
checa(19, "baixar o relatorio de obrigacoes emite EXPORT_DADOS",
      e_export is not None)
checa(20, f"e a linha diz QUANTAS linhas sairam, e o numero bate com o cadastro "
          f"(saiu: {e_export.get('quantidade') if e_export else None})",
      e_export is not None and e_export.get("quantidade") == 2
      and e_export.get("recurso") == "obrigacoes")

# O modelo de importacao e uma planilha VAZIA de layout, e nao dado: marcar o
# download dele como export encheria o log de falso positivo.
_, l_modelo = registra(
    lambda: cliente.get("/api/obrigacoes/modelo-importacao", headers=admin))
checa(21, "baixar o MODELO em branco nao e export de dado, e nao emite nada",
      evento(l_modelo, "EXPORT_DADOS") is None)


# ------------------------------------ (h) nomes padronizados e lista proibida
NOMES_PROIBIDOS = ["DOCUMENTO_EXCLUIDO"]
fora_da_tabela = []
for linha in todas_as_linhas:
    if linha.get("event") in NOMES_PROIBIDOS:
        fora_da_tabela.append(linha.get("event"))
checa(22, f"nenhum nome de evento fora da tabela da nota (achados: "
          f"{sorted(set(fora_da_tabela))})",
      not fora_da_tabela)

novos = [linha for linha in todas_as_linhas
         if linha.get("event") in (CRIACAO, EDICAO, EXCLUSAO, "EXPORT_DADOS")]
checa(23, f"toda linha nova traz os oito campos na ordem da tabela "
          f"(linhas: {len(novos)})",
      len(novos) >= 10 and all(list(l)[:8] == CAMPOS for l in novos))

achados = []
for linha in novos:
    for chave in linha:
        for proibida in PROIBIDAS:
            if proibida in chave.lower():
                achados.append((linha.get("event"), chave))
checa(24, f"nenhuma linha nova carrega campo da lista proibida "
          f"(achados: {achados})",
      not achados)


# -------------------- (i) os caminhos laterais, achado do verificador
# Cobrir a rota que o levantamento apontou e nao procurar o PADRAO IRMAO e o
# erro que a `Escada_Preguica_de_Codigo` nomeia, e ele se repetiu aqui: sete
# rotas mudam as mesmas quatro tabelas criticas por outro caminho e nao
# deixavam rastro nenhum. A mais gritante e o `status`, que mexe no MESMO
# campo `ativa` que o soft delete mexe e que ja registrava.
db = SessionLocal()
emp2 = Empresa(razao_social="Empresa do Bloqueio", cnpj="19131243000197")
usr2 = Usuario(nome="Bloqueavel", email="bloqueavel@bps4.com.br",
               grupo="analista", senha_hash=get_password_hash(SENHA),
               ativo=True)
obr2 = Obrigacao(nome="GIA mensal", setor_id=ids["setor"])
db.add_all([emp2, usr2, obr2])
db.commit()
ids.update(emp2=emp2.id, usr2=usr2.id, obr2=obr2.id)
db.close()

_, l_bloq_emp = registra(
    lambda: cliente.post(f"/api/empresas/{ids['emp2']}/bloquear",
                         json={"bloqueado": True}, headers=admin))
checa(25, "bloquear empresa emite EDICAO: e mudanca de estado critico",
      evento(l_bloq_emp, EDICAO) is not None)

_, l_bloq_usr = registra(
    lambda: cliente.post(f"/api/usuarios/{ids['usr2']}/bloquear",
                         json={"bloqueado": True}, headers=admin))
checa(26, "bloquear usuario emite EDICAO: mexe em quem consegue entrar",
      evento(l_bloq_usr, EDICAO) is not None)

# Desativar pelo `status` tem o MESMO efeito que o soft delete do DELETE, que
# ja emitia EXCLUSAO com `inativado=True`. Sai com o mesmo evento, senao quem
# filtrar exclusao acha um caminho e perde o outro.
_, l_status_off = registra(
    lambda: cliente.post(f"/api/obrigacoes/{ids['obr2']}/status",
                         json={"ativa": False}, headers=admin))
e_status = evento(l_status_off, EXCLUSAO)
checa(27, "desativar obrigacao pelo status emite EXCLUSAO com inativado=True, "
          "igual ao soft delete que mexe no mesmo campo",
      e_status is not None and e_status.get("inativado") is True)

_, l_status_on = registra(
    lambda: cliente.post(f"/api/obrigacoes/{ids['obr2']}/status",
                         json={"ativa": True}, headers=admin))
checa(28, "e reativar emite EDICAO, porque ressuscitar nao e excluir",
      evento(l_status_on, EDICAO) is not None
      and evento(l_status_on, EXCLUSAO) is None)

_, l_copiar = registra(
    lambda: cliente.post("/api/obrigacoes/copiar-empresa",
                         json={"origem_empresa_id": ids["empresa"],
                               "destino_empresa_id": ids["emp2"]},
                         headers=admin))
checa(29, "copiar obrigacoes de uma empresa para outra emite UMA linha de EDICAO",
      len(eventos(l_copiar, EDICAO)) == 1)

_, l_desvincula = registra(
    lambda: cliente.post("/api/obrigacoes/desvincular-empresa",
                         json={"empresa_id": ids["emp2"]}, headers=admin))
checa(30, "desvincular empresa das obrigacoes emite UMA linha de EDICAO",
      len(eventos(l_desvincula, EDICAO)) == 1)

_, l_detalhe = registra(
    lambda: cliente.put(f"/api/obrigacoes/{ids['obr2']}/detalhes-empresa",
                        json={"itens": [{"empresa_id": ids["empresa"],
                                         "observacao": "conta do Banco X"}]},
                        headers=admin))
checa(31, "regravar detalhes por empresa emite EDICAO, como a rota irma de "
          "responsaveis por setor ja emitia",
      evento(l_detalhe, EDICAO) is not None)


def _planilha_usuarios():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["nome", "email", "nivel"])
    ws.append(["Importado Um", "imp1@bps4.com.br", "analista"])
    ws.append(["Importado Dois", "imp2@bps4.com.br", "analista"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_, l_imp_usr = registra(
    lambda: cliente.post(
        "/api/usuarios/importar",
        files={"arquivo": ("usuarios.xlsx", _planilha_usuarios(),
                           "application/vnd.openxmlformats-officedocument"
                           ".spreadsheetml.sheet")},
        headers=admin))
criacoes_usr = eventos(l_imp_usr, CRIACAO)
checa(32, f"importar usuarios emite UMA linha com a contagem, como a de "
          f"empresas ja emitia (saiu: {len(criacoes_usr)})",
      len(criacoes_usr) == 1
      and criacoes_usr[0].get("lote") is True
      and criacoes_usr[0].get("quantidade") == 2)


# --------- (j) linha de log nao pode afirmar o que o banco nao gravou
# Achado adversarial: o `MUDANCA_ROLE` saia ANTES do `db.commit()`. Se o commit
# estoura, a linha ja anunciou uma promocao de privilegio que nunca aconteceu,
# e log de auditoria que mente sobre elevacao de acesso e pior do que log
# nenhum: manda quem investiga para o lado errado com ar de prova.
# O caminho de falha nao e inventado: o PUT de usuario NAO confere e-mail
# duplicado (so a criacao confere), entao trocar papel e e-mail no mesmo pedido
# faz o banco recusar por unique constraint.
_, l_rollback = registra(
    lambda: cliente_erro.put(f"/api/usuarios/{ids['usr2']}",
                        json={"grupo": "admin", "email": EMAIL_ADMIN},
                        headers=admin))

db = SessionLocal()
papel_no_banco = (db.query(Usuario)
                  .filter(Usuario.id == ids["usr2"]).first().grupo)
db.close()

checa(33, f"commit que falha NAO deixa linha de MUDANCA_ROLE "
          f"(papel no banco: {papel_no_banco})",
      papel_no_banco == "analista"
      and evento(l_rollback, "MUDANCA_ROLE") is None)
checa(34, "e tambem nao deixa linha de EDICAO afirmando o que nao gravou",
      evento(l_rollback, EDICAO) is None)


# O mesmo padrao no importador, onde o commit e unico e roda depois do laco
# inteiro: uma falha na linha N descartaria as N-1 anteriores, e as linhas de
# log delas ja teriam saido. Forcado com o `commit` real quebrando.
from sqlalchemy.orm import Session as _Sessao                   # noqa: E402

_commit_real = _Sessao.commit


def _commit_que_estoura(self):
    raise RuntimeError("commit derrubado de proposito pela prova")


def _planilha_troca_papel():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["nome", "email", "nivel"])
    ws.append(["Importado Um", "imp1@bps4.com.br", "admin"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_Sessao.commit = _commit_que_estoura
try:
    _, l_imp_quebrado = capturar(
        lambda: cliente_erro.post(
            "/api/usuarios/importar",
            files={"arquivo": ("usuarios.xlsx", _planilha_troca_papel(),
                               "application/vnd.openxmlformats-officedocument"
                               ".spreadsheetml.sheet")},
            headers=admin))
finally:
    _Sessao.commit = _commit_real

checa(35, "importacao cujo commit estoura NAO deixa linha de MUDANCA_ROLE",
      evento(l_imp_quebrado, "MUDANCA_ROLE") is None)
checa(36, "nem linha de CRIACAO em lote",
      evento(l_imp_quebrado, CRIACAO) is None)


# --------------- (k) lote que nao teve efeito nenhum nao anuncia criacao
# A planilha sem a coluna obrigatoria devolve 200 com um erro dentro, e criava
# zero. A linha saia assim mesmo, dizendo `quantidade: 0`: o numero nao mente,
# mas o EVENTO mente, porque anuncia criacao critica onde nao houve criacao
# nenhuma. As rotas singulares nunca logam em caminho de erro, e o lote passa a
# seguir a mesma regra.
def _planilha_sem_coluna():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["coluna estranha", "outra"])
    ws.append(["nada", "aqui"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_, l_lote_vazio = registra(
    lambda: cliente.post(
        "/api/empresas/importar",
        files={"arquivo": ("empresas.xlsx", _planilha_sem_coluna(),
                           "application/vnd.openxmlformats-officedocument"
                           ".spreadsheetml.sheet")},
        headers=admin))
checa(37, "importacao que nao criou nem atualizou nada NAO emite evento",
      evento(l_lote_vazio, CRIACAO) is None)


print()
if falhou:
    print(f"PROVA FALHOU nos itens: {sorted(set(falhou))}")
    sys.exit(1)
print(f"PROVA OK: {37 - len(falhou)} checagens verdes")
