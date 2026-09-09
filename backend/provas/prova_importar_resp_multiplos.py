"""
prova_importar_resp_multiplos.py — a planilha de responsáveis aceita mais de um
nome por célula.

O escritório cadastra a matriz inteira por Excel, e não linha a linha na tela.
Se a planilha continuasse aceitando um nome só, ela seria o caminho que DESFAZ
o que a tela permite: reimportar apagaria o segundo responsável de cada setor
sem dizer nada.

O que se prova aqui:
· "Ana; Bruno" na mesma célula grava os dois, nessa ordem, com Ana de principal;
· um nome desconhecido no meio vira aviso e NÃO derruba a linha, nem os nomes
  que casaram;
· célula vazia continua desmarcando o setor, como sempre foi;
· o modelo baixável ensina o separador, senão o recurso existe sem ninguém usar.

Rodar:  cd backend && ./venv/bin/python provas/prova_importar_resp_multiplos.py
"""
import io
import os
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
os.environ.setdefault("DATABASE_URL", "sqlite:///" + tempfile.mktemp(suffix=".db"))

from app.database import SessionLocal, Base, engine                   # noqa: E402
from app.models import (Empresa, Setor, Usuario,                      # noqa: E402
                        EmpresaSetorResponsavel,
                        empresa_setor_resp_usuarios)
from app.services import importador_resp_setor as imp                 # noqa: E402

ok = True


def check(nome, cond, extra=""):
    global ok
    print(("  OK   " if cond else "  FALHA ") + nome + (f"  {extra}" if extra else ""))
    ok = ok and bool(cond)


Base.metadata.create_all(bind=engine)


def planilha(cabecalho, linhas):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(cabecalho)
    for l in linhas:
        ws.append(l)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def montar():
    db = SessionLocal()
    db.execute(empresa_setor_resp_usuarios.delete())
    for m in (EmpresaSetorResponsavel, Empresa, Setor, Usuario):
        db.query(m).delete()
    db.commit()
    db.add(Empresa(razao_social="ACME", cnpj="11.222.333/0001-81", ativo=True))
    for n in ("Fiscal", "Contábil"):
        db.add(Setor(nome=n, ativo=True))
    for n in ("Ana Paula", "Bruno Sá", "Carla"):
        db.add(Usuario(nome=n, email=f"{n.split()[0].lower()}@x.com",
                       senha_hash="x", grupo="analista", tipo="colaborador"))
    db.commit()
    return db


def lista_do_setor(db, nome_setor):
    s = db.query(Setor).filter(Setor.nome == nome_setor).first()
    v = (db.query(EmpresaSetorResponsavel)
         .filter(EmpresaSetorResponsavel.setor_id == s.id).first())
    if not v:
        return None
    return [u.nome for u in v.responsaveis], v.responsavel_id


print("\n=== 1. dois nomes na mesma célula ===")
db = montar()
r = imp.importar(db, "resp.xlsx",
                 planilha(["CNPJ", "Fiscal", "Contábil"],
                          [["11222333000181", "Ana Paula; Bruno Sá", "Carla"]]))
nomes, principal = lista_do_setor(db, "Fiscal")
check("os dois entraram, na ordem da célula", nomes == ["Ana Paula", "Bruno Sá"],
      f"({nomes})")
ana = db.query(Usuario).filter(Usuario.nome == "Ana Paula").first()
check("a primeira da célula é a principal", principal == ana.id)
check("a outra coluna, com um nome só, continua funcionando",
      lista_do_setor(db, "Contábil")[0] == ["Carla"])
check("a importação não acusou erro", r["resumo"]["erros"] == 0, f"({r['resumo']})")
db.close()

print("\n=== 2. nome desconhecido no meio não derruba a linha ===")
db = montar()
r = imp.importar(db, "resp.xlsx",
                 planilha(["CNPJ", "Fiscal"],
                          [["11222333000181", "Ana Paula; Fulano Que Nao Existe; Carla"]]))
nomes, _ = lista_do_setor(db, "Fiscal")
check("os nomes que casaram entraram assim mesmo", nomes == ["Ana Paula", "Carla"],
      f"({nomes})")
avisos = [d for d in r["detalhes"] if d["status"] == "aviso"]
check("saiu UM aviso, dizendo qual nome ficou de fora", len(avisos) == 1
      and "Fulano Que Nao Existe" in avisos[0]["detalhe"], f"({avisos})")
check("aviso não é erro: a linha foi importada", r["resumo"]["erros"] == 0
      and r["resumo"]["marcados"] == 1, f"({r['resumo']})")
db.close()

print("\n=== 3. nenhum nome da célula existe ===")
db = montar()
r = imp.importar(db, "resp.xlsx",
                 planilha(["CNPJ", "Fiscal"], [["11222333000181", "Zé; Zico"]]))
nomes, principal = lista_do_setor(db, "Fiscal")
check("o setor fica marcado, e sem responsável", nomes == [] and principal is None,
      f"({nomes}, {principal})")
check("o aviso diz que ninguém foi encontrado",
      any("sem responsável" in d["detalhe"] for d in r["detalhes"]), f"({r['detalhes']})")
db.close()

print("\n=== 4. célula vazia continua desmarcando o setor ===")
db = montar()
imp.importar(db, "resp.xlsx",
             planilha(["CNPJ", "Fiscal", "Contábil"],
                      [["11222333000181", "Ana Paula; Bruno Sá", "Carla"]]))
r = imp.importar(db, "resp.xlsx",
                 planilha(["CNPJ", "Fiscal", "Contábil"],
                          [["11222333000181", "", "Carla"]]))
check("o setor com célula vazia sai da matriz", lista_do_setor(db, "Fiscal") is None)
check("e o outro continua lá", lista_do_setor(db, "Contábil")[0] == ["Carla"])
check("a contagem de desmarcados subiu", r["resumo"]["desmarcados"] == 1,
      f"({r['resumo']})")
check("desmarcar não deixa linha órfã na tabela do meio",
      db.query(empresa_setor_resp_usuarios).count() == 1,
      f"({db.query(empresa_setor_resp_usuarios).count()} linhas)")
db.close()

print("\n=== 5. reimportar troca a lista inteira, sem empilhar ===")
db = montar()
imp.importar(db, "resp.xlsx",
             planilha(["CNPJ", "Fiscal"], [["11222333000181", "Ana Paula; Bruno Sá"]]))
imp.importar(db, "resp.xlsx",
             planilha(["CNPJ", "Fiscal"], [["11222333000181", "Carla"]]))
nomes, _ = lista_do_setor(db, "Fiscal")
check("a segunda importação substitui, não soma", nomes == ["Carla"], f"({nomes})")
check("e não sobrou linha da lista anterior",
      db.query(empresa_setor_resp_usuarios).count() == 1,
      f"({db.query(empresa_setor_resp_usuarios).count()} linhas)")
db.close()

print("\n=== 6. o nome repetido na célula colapsa ===")
db = montar()
imp.importar(db, "resp.xlsx",
             planilha(["CNPJ", "Fiscal"], [["11222333000181", "Ana Paula; ana paula"]]))
nomes, _ = lista_do_setor(db, "Fiscal")
check("a mesma pessoa duas vezes vira uma", nomes == ["Ana Paula"], f"({nomes})")
db.close()

print("\n=== 7. o modelo baixável ensina o separador ===")
db = montar()
import openpyxl                                                        # noqa: E402
ws = openpyxl.load_workbook(io.BytesIO(imp.gerar_modelo(db))).worksheets[0]
texto = "\n".join(str(c.value or "") for l in ws.iter_rows() for c in l)
check("o exemplo mostra dois nomes separados por ponto e vírgula",
      "Ana Paula; Bruno Sá" in texto)
check("e a instrução explica a regra por extenso",
      "ponto e vírgula" in texto and "principal" in texto)
db.close()

print("\n" + ("TODAS AS PROVAS PASSARAM" if ok else "HOUVE FALHA"))
sys.exit(0 if ok else 1)
